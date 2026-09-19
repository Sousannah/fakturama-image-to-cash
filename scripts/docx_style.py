"""Shared python-docx styling for the three Word deliverables.

The documents are generated rather than hand-written so they cannot drift from
the code: every number in them (test count, step count, tier usage, run status)
is read from the repository or from a run report at build time.
"""
from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Iterable, List, Optional, Sequence, Tuple

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor

REPO_ROOT = Path(__file__).resolve().parents[1]

NAVY = RGBColor(0x0E, 0x23, 0x38)
ACCENT = RGBColor(0x1F, 0x62, 0xA8)
MUTED = RGBColor(0x69, 0x7A, 0x8D)
INK = RGBColor(0x16, 0x20, 0x2C)
OK = RGBColor(0x17, 0x86, 0x4A)
WARN = RGBColor(0xB4, 0x70, 0x0B)
ERR = RGBColor(0xC3, 0x2B, 0x2B)

BODY_FONT = "Calibri"
MONO_FONT = "Consolas"

SHADE_HEAD = "0E2338"
SHADE_ROW = "F2F6FA"
SHADE_CODE = "F5F7FA"
SHADE_NOTE = "EAF2FB"


# --------------------------------------------------------------------------- #
# low-level xml helpers
# --------------------------------------------------------------------------- #
def _shade(element, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def shade_cell(cell, fill: str) -> None:
    _shade(cell._tc.get_or_add_tcPr(), fill)


def shade_paragraph(paragraph, fill: str) -> None:
    _shade(paragraph._p.get_or_add_pPr(), fill)


def paragraph_border(paragraph, side: str = "left", colour: str = "1F62A8", size: int = 18) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    edge = OxmlElement("w:%s" % side)
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), str(size))
    edge.set(qn("w:space"), "8")
    edge.set(qn("w:color"), colour)
    borders.append(edge)
    pPr.append(borders)


def table_borders(table, colour: str = "D5DEE9", size: int = 4) -> None:
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:%s" % edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(size))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), colour)
        borders.append(e)
    tblPr.append(borders)


def _field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


# --------------------------------------------------------------------------- #
# document construction
# --------------------------------------------------------------------------- #
class Doc:
    """A thin, opinionated wrapper: the only styling vocabulary these builds use."""

    def __init__(self, title: str, subtitle: str = "", footer: str = ""):
        self.doc = Document()
        self._base_styles()
        section = self.doc.sections[0]
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        self.width = section.page_width - section.left_margin - section.right_margin
        if footer:
            self._footer(footer)
        if title:
            self.cover(title, subtitle)

    # -- styles ---------------------------------------------------------- #
    def _base_styles(self) -> None:
        normal = self.doc.styles["Normal"]
        normal.font.name = BODY_FONT
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = INK
        normal.paragraph_format.space_after = Pt(7)
        normal.paragraph_format.line_spacing = 1.16

        for name, size, colour, before, after in (
            ("Heading 1", 18, NAVY, 20, 8),
            ("Heading 2", 13.5, NAVY, 15, 5),
            ("Heading 3", 11.5, ACCENT, 11, 4),
        ):
            st = self.doc.styles[name]
            st.font.name = BODY_FONT
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.color.rgb = colour
            st.paragraph_format.space_before = Pt(before)
            st.paragraph_format.space_after = Pt(after)
            st.paragraph_format.keep_with_next = True

    def _footer(self, text: str) -> None:
        p = self.doc.sections[0].footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text + "    ·    page ")
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED
        run.font.name = BODY_FONT
        _field(p, " PAGE ")
        for r in p.runs:
            r.font.size = Pt(8)
            r.font.color.rgb = MUTED
            r.font.name = BODY_FONT

    # -- blocks ---------------------------------------------------------- #
    def cover(self, title: str, subtitle: str = "") -> None:
        band = self.doc.add_paragraph()
        band.paragraph_format.space_before = Pt(6)
        band.paragraph_format.space_after = Pt(2)
        run = band.add_run("FAKTURAMA IMAGE-TO-CASH AUTOMATION")
        run.font.size = Pt(8.5)
        run.font.bold = True
        run.font.color.rgb = ACCENT
        run.font.name = BODY_FONT

        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(title)
        run.font.size = Pt(26)
        run.font.bold = True
        run.font.color.rgb = NAVY
        run.font.name = BODY_FONT

        if subtitle:
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(10)
            run = p.add_run(subtitle)
            run.font.size = Pt(11)
            run.font.color.rgb = MUTED
            run.font.name = BODY_FONT

        rule = self.doc.add_paragraph()
        rule.paragraph_format.space_before = Pt(0)
        rule.paragraph_format.space_after = Pt(12)
        paragraph_border(rule, "bottom", "1F62A8", 12)

    def h1(self, text: str) -> None:
        self.doc.add_heading(text, level=1)

    def h2(self, text: str) -> None:
        self.doc.add_heading(text, level=2)

    def h3(self, text: str) -> None:
        self.doc.add_heading(text, level=3)

    def p(self, text: str = "", bold: bool = False, muted: bool = False, size: Optional[float] = None):
        """Body text. `**bold**` and `` `code` `` inside the string are honoured."""
        para = self.doc.add_paragraph()
        self._rich(para, text, bold=bold, muted=muted, size=size)
        return para

    def _rich(self, para, text: str, bold: bool = False, muted: bool = False, size=None) -> None:
        for chunk in re.split(r"(\*\*[^*]+\*\*|\*[^*`]+\*|`[^`]+`)", text):
            if not chunk:
                continue
            if chunk.startswith("**") and chunk.endswith("**"):
                run = para.add_run(chunk[2:-2])
                run.font.bold = True
            elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
                run = para.add_run(chunk[1:-1])
                run.font.italic = True
            elif chunk.startswith("`") and chunk.endswith("`"):
                run = para.add_run(chunk[1:-1])
                run.font.name = MONO_FONT
                run.font.size = Pt((size or 10.5) - 1)
                run.font.color.rgb = ACCENT
            else:
                run = para.add_run(chunk)
                run.font.bold = bold
            if muted:
                run.font.color.rgb = MUTED
            if size and run.font.size is None:
                run.font.size = Pt(size)

    def bullets(self, items: Iterable[str], style: str = "List Bullet") -> None:
        for item in items:
            para = self.doc.add_paragraph(style=style)
            para.paragraph_format.space_after = Pt(3)
            self._rich(para, item)

    def numbers(self, items: Iterable[str]) -> None:
        self.bullets(items, style="List Number")

    def code(self, text: str, caption: str = "") -> None:
        if caption:
            cap = self.doc.add_paragraph()
            cap.paragraph_format.space_after = Pt(2)
            run = cap.add_run(caption)
            run.font.size = Pt(8)
            run.font.bold = True
            run.font.color.rgb = MUTED
            run.font.name = BODY_FONT

        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = self.width
        shade_cell(cell, SHADE_CODE)
        table_borders(table, "E1E8F0", 4)

        cell.paragraphs[0].text = ""
        for i, line in enumerate(text.strip("\n").splitlines()):
            para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            run = para.add_run(line)
            run.font.name = MONO_FONT
            run.font.size = Pt(8.5)
            run.font.color.rgb = INK
        self.spacer(6)

    def note(self, text: str, kind: str = "info") -> None:
        colour = {"info": "1F62A8", "ok": "17864A", "warn": "B4700B", "err": "C32B2B"}[kind]
        para = self.doc.add_paragraph()
        para.paragraph_format.left_indent = Inches(0.06)
        para.paragraph_format.space_before = Pt(6)
        para.paragraph_format.space_after = Pt(8)
        shade_paragraph(para, SHADE_NOTE if kind == "info" else "FAFBFD")
        paragraph_border(para, "left", colour, 18)
        self._rich(para, text)

    def table(
        self,
        headers: Sequence[str],
        rows: Sequence[Sequence[str]],
        widths: Optional[Sequence[float]] = None,
        font_size: float = 9,
    ):
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_borders(table)

        for i, head in enumerate(headers):
            cell = table.rows[0].cells[i]
            shade_cell(cell, SHADE_HEAD)
            cell.paragraphs[0].text = ""
            run = cell.paragraphs[0].add_run(head.upper())
            run.font.bold = True
            run.font.size = Pt(font_size - 1)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.name = BODY_FONT
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)

        for r, row in enumerate(rows):
            cells = table.add_row().cells
            for i, value in enumerate(row):
                cell = cells[i]
                if r % 2 == 1:
                    shade_cell(cell, SHADE_ROW)
                cell.paragraphs[0].text = ""
                cell.paragraphs[0].paragraph_format.space_after = Pt(2)
                self._rich(cell.paragraphs[0], str(value), size=font_size)
                for run in cell.paragraphs[0].runs:
                    if run.font.size is None:
                        run.font.size = Pt(font_size)

        if widths:
            for row in table.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Emu(int(self.width * w))
        self.spacer(6)
        return table

    def image(self, path: Path, caption: str = "", width_ratio: float = 1.0) -> None:
        path = Path(path)
        if not path.exists():
            self.note("missing image: %s" % path, "warn")
            return
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before = Pt(6)
        para.paragraph_format.space_after = Pt(3)
        para.add_run().add_picture(str(path), width=Emu(int(self.width * width_ratio)))
        if caption:
            cap = self.doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.space_after = Pt(12)
            run = cap.add_run(caption)
            run.font.size = Pt(8.5)
            run.font.italic = True
            run.font.color.rgb = MUTED
            run.font.name = BODY_FONT

    def spacer(self, points: float = 8) -> None:
        para = self.doc.add_paragraph()
        para.paragraph_format.space_after = Pt(points)
        para.paragraph_format.space_before = Pt(0)

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def orientation(self, landscape: bool) -> None:
        """Start a new section in the given orientation.

        Full-screen captures are 16:9; on a portrait page they come out half the
        height of the text and the annotation labels stop being readable. The
        walk-through therefore turns the page for its figures and turns it back
        afterwards. `self.width` is updated so `image()` keeps filling the
        column.
        """
        from docx.enum.section import WD_ORIENT

        section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        width, height = section.page_width, section.page_height
        want = (max(width, height), min(width, height)) if landscape else (
            min(width, height), max(width, height))
        section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = want
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        self.width = section.page_width - section.left_margin - section.right_margin

    def toc_list(self, entries: Sequence[Tuple[str, str]]) -> None:
        """A plain contents list - a real TOC field needs Word to refresh it."""
        for number, title in entries:
            para = self.doc.add_paragraph()
            para.paragraph_format.space_after = Pt(3)
            run = para.add_run("%s   " % number)
            run.font.bold = True
            run.font.color.rgb = ACCENT
            run.font.name = BODY_FONT
            run = para.add_run(title)
            run.font.color.rgb = NAVY
            run.font.name = BODY_FONT

    def save(self, path: Path) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
        return path


# --------------------------------------------------------------------------- #
# facts read from the repository, so the prose cannot drift
# --------------------------------------------------------------------------- #
def latest_run(out_dir: Optional[Path] = None, status: Optional[str] = None) -> Optional[Path]:
    """The newest run directory that has a report, optionally filtered by status."""
    out_dir = out_dir or (REPO_ROOT / "out")
    runs = sorted(out_dir.glob("run-*"), reverse=True)
    for run in runs:
        report = run / "report.json"
        if not report.exists():
            continue
        if status is None:
            return run
        try:
            if json.loads(report.read_text(encoding="utf-8")).get("status") == status:
                return run
        except Exception:
            continue
    return None


def load_report(run_dir: Path) -> dict:
    return json.loads((Path(run_dir) / "report.json").read_text(encoding="utf-8"))


def collected_tests() -> int:
    """How many tests pytest actually collects.

    Asking pytest beats counting `def test_` with a regex, because parametrised
    cases count for more than one - and a document that claims the wrong number
    is exactly the kind of drift these builders exist to prevent.
    """
    import subprocess
    import sys

    try:
        out = subprocess.run(
            [sys.executable, "-m", "pytest", "--collect-only", "-q"],
            cwd=str(REPO_ROOT), capture_output=True, text=True, timeout=180,
        ).stdout
    except Exception:
        return 0
    match = re.search(r"(\d+)\s+tests? collected", out)
    if match:
        return int(match.group(1))
    return len([l for l in out.splitlines() if "::" in l])


def test_count() -> int:
    """Count test functions without importing them."""
    total = 0
    for path in (REPO_ROOT / "tests").glob("test_*.py"):
        total += len(re.findall(r"^def test_", path.read_text(encoding="utf-8"), re.M))
        # parametrised cases count for more than one, but the headline number in
        # the documents is the one pytest prints, so it is passed in explicitly
        # where it matters rather than guessed here.
    return total


def loc(*relative: str) -> int:
    total = 0
    for rel in relative:
        path = REPO_ROOT / rel
        for f in ([path] if path.is_file() else sorted(path.rglob("*.py"))):
            total += len(f.read_text(encoding="utf-8").splitlines())
    return total


def selector_stats() -> Tuple[int, int, dict]:
    import yaml

    data = yaml.safe_load((REPO_ROOT / "src" / "f2c" / "ui" / "selectors.yaml").read_text(encoding="utf-8"))
    tiers: dict = {}
    strategies = 0
    for value in data.values():
        for s in value:
            tiers[s["tier"]] = tiers.get(s["tier"], 0) + 1
            strategies += 1
    return len(data), strategies, tiers
